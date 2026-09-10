"""Excel de sanciones — trasplante de `nucleo_modular/reportes.generar_excel_sanciones`
(origen `procesador.ProcesadorRRHH.exportar_a_excel`). Funciones puras: datos -> bytes.

Una hoja por categoría + "Detalle Resumen" (con valor monetario por tipo) + "Resumen".
Se conserva openpyxl (multi-hoja con estilos, puerto directo del legado).
"""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO

from core.repos.sanciones import categorizar_sanciones

log = logging.getLogger(__name__)

_HEADERS = [
    "ID", "Empleado Cod", "Empleado Cédula", "Empleado Nombre", "Puesto", "Agente",
    "Fecha", "Hora", "Tipo Sanción", "Observaciones", "Observaciones Adicionales",
    "Horas Extras", "Status", "Comentarios Gerencia", "Comentarios RRHH",
    "Pendiente", "Foto URL", "Firma Path",
    "Supervisor ID", "Nombre Supervisor", "Fecha Revisión",
    "Reviewed By", "Nombre Revisor", "Created At", "Updated At",
    "Procesado Por", "Fecha Procesamiento",
]
_WIDTHS = [12, 12, 15, 25, 20, 15, 12, 10, 18, 40, 40, 12, 12, 40, 40,
          10, 25, 25, 15, 30, 18, 15, 25, 18, 18, 15, 18]


def _safe(v: object) -> object:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "SÍ" if v else "NO"
    if isinstance(v, (int, float, str)):
        return v
    return str(v)


def _ced(cedula: object) -> str:
    if not cedula or str(cedula).strip() in ("", "N/A", "None"):
        return ""
    return str(cedula).strip()


def _parse_fecha(v: object):
    if not v:
        return None
    if hasattr(v, "strftime"):
        return v
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(v)[:10], fmt)
        except ValueError:
            continue
    return None


def _valor_multa(sancion: dict, valores: dict[str, float]) -> object:
    """Fórmula de la hoja "Detalle Resumen" del legado (procesador.py:1906-1935)."""
    tipo = (sancion.get("tipo_sancion") or "").strip().upper()
    obs = (sancion.get("observaciones") or "").strip().lower()
    if tipo in ("FALTA", "PERMISO"):
        return valores.get(tipo, 1)
    if tipo in ("HORAS EXTRAS", "MANIOBRAS"):
        if "12 h" in obs or "12 horas" in obs:
            return valores.get("HORAS EXTRAS 12H (FIJO)", 30)
        factor = valores.get("HORAS EXTRAS (FACTOR/HORA)", 2.5)
        try:
            he = float(sancion.get("horas_extras") or 0)
            return round(he * factor, 2) if he else ""
        except (TypeError, ValueError):
            return ""
    if tipo == "FRANCO TRABAJADO":
        if "12 h" in obs or "12 horas" in obs:
            return valores.get("FRANCO TRABAJADO 12H", 37.5)
        return valores.get("FRANCO TRABAJADO 8H", 25)
    return valores.get(tipo, "")


def sanciones_xlsx(
    sanciones: list[dict],
    nombres_supervisores: dict[str, str] | None = None,
    valores_sanciones: dict[str, float] | None = None,
) -> bytes | None:
    """`.xlsx` de sanciones ya enriquecidas. `None` si no hay sanciones o falta openpyxl."""
    if not sanciones:
        return None
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.error("openpyxl no está instalado")
        return None

    nombres = nombres_supervisores or {}
    usar_nombres = bool(nombres)
    valores = valores_sanciones or {}
    categorizadas = categorizar_sanciones(sanciones)

    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    halign = Alignment(horizontal="center", vertical="center")
    hojas = 0

    def _nom_sup(sid: str) -> str:
        if sid and usar_nombres:
            return nombres.get(sid, f"No encontrado ({sid[:8]}...)" if sid else "")
        return "Sin acceso a profiles" if sid else ""

    for categoria, filas in categorizadas.items():
        if not filas:
            continue
        ws = wb.create_sheet(title=categoria.replace("/", "-").replace("\\", "-")[:30])
        for c, h in enumerate(_HEADERS, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.alignment = hfont, hfill, halign
        for r, s in enumerate(filas, 2):
            row = [
                _safe(s.get("id", "")), _safe(s.get("empleado_cod", "")), _ced(s.get("empleado_cedula")),
                _safe(s.get("empleado_nombre", "")), _safe(s.get("puesto", "")), _safe(s.get("agente", "")),
                _safe(s.get("fecha", "")), _safe(s.get("hora", "")), _safe(s.get("tipo_sancion", "")),
                _safe(s.get("observaciones", "")), _safe(s.get("observaciones_adicionales", "")),
                _safe(s.get("horas_extras", "")), _safe(s.get("status", "")),
                _safe(s.get("comentarios_gerencia", "")), _safe(s.get("comentarios_rrhh", "")),
                _safe(s.get("pendiente", "")), _safe(s.get("foto_url", "")), _safe(s.get("firma_path", "")),
                _safe(s.get("supervisor_id", "")), _safe(_nom_sup(s.get("supervisor_id", ""))),
                _safe(s.get("fecha_revision", "")), _safe(s.get("reviewed_by", "")),
                _safe(_nom_sup(s.get("reviewed_by", ""))), _safe(s.get("created_at", "")),
                _safe(s.get("updated_at", "")), _safe(s.get("procesado_por", "")),
                _safe(s.get("fecha_procesamiento", "")),
            ]
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for c, w in enumerate(_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        for r in range(2, ws.max_row + 1):
            ws[f"{get_column_letter(3)}{r}"].number_format = "@"
        hojas += 1

    if hojas:
        ws = wb.create_sheet(title="Detalle Resumen")
        dfont = Font(bold=True, color="FFFFFF")
        dfill = PatternFill(start_color="A23B72", end_color="A23B72", fill_type="solid")
        for c, h in enumerate(
            ["ID", "Empleado Cod", "Empleado Cédula", "Empleado Nombre", "Tipo Sanción",
             "Nombre Supervisor", "Fecha (dd/mm/aaaa)", "Valor", "Observacion"], 1,
        ):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.alignment = dfont, dfill, halign

        def _clean(v: object) -> str:
            return str(v).strip() if v else ""

        for r, s in enumerate(sanciones, 2):
            nom_sup = nombres.get(s.get("supervisor_id", ""), "") if usar_nombres else ""
            fdt = _parse_fecha(s.get("fecha"))
            fecha_fmt = fdt.strftime("%d/%m/%Y") if fdt else ""
            tipo = _clean(s.get("tipo_sancion"))
            partes = [p for p in [
                f"{_clean(s.get('id', ''))[:4]}-SANCION", fecha_fmt, tipo,
                _clean(s.get("observaciones", ""))[:50], _clean(s.get("observaciones_adicionales", "")),
                _clean(s.get("comentarios_gerencia", "")), _clean(nom_sup),
            ] if p]
            row = [
                _safe(s.get("id", "")), _safe(s.get("empleado_cod", "")), _ced(s.get("empleado_cedula")),
                _safe(s.get("empleado_nombre", "")), tipo, _safe(nom_sup), fecha_fmt,
                _valor_multa(s, valores), " ".join(partes),
            ]
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for c, w in enumerate([36, 14, 15, 28, 28, 28, 16, 8, 110], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        for r in range(2, ws.max_row + 1):
            ws[f"{get_column_letter(3)}{r}"].number_format = "@"
        ws.freeze_panes = "A2"

        wr = wb.create_sheet(title="Resumen", index=0)
        wr["A1"] = "RESUMEN DE EXPORTACIÓN"
        wr["A3"] = f"Fecha de generación: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        wr["A4"] = f"Total de sanciones: {len(sanciones)}"
        wr["A5"] = f"Hojas creadas: {hojas}"
        wr["A6"] = f"Nombres de supervisores: {'Incluidos' if usar_nombres else 'No disponibles'}"
        fila = 8
        wr[f"A{fila}"] = "DETALLE POR CATEGORÍA:"
        for categoria, filas in categorizadas.items():
            if filas:
                fila += 1
                wr[f"A{fila}"] = f"• {categoria}: {len(filas)} sanciones"
        wr.column_dimensions["A"].width = 40

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
