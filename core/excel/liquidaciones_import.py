"""
Importador de liquidaciones ya calculadas desde un Excel externo (formato
LIQUIDACIONES_REG de 62+ columnas del `.pyw` legado, el que genera
`Generador_Liquidaciones_INSEVIG.pyw`) -- para cargas históricas donde se
decide respetar los montos ya calculados/validados por RRHH en vez de
recalcular contra RPEMPLEA/movimientos actuales (que pueden haber cambiado
desde que se generó el Excel -- ver docs/modulos/liquidaciones.md, caso
real: "Liquidaciones_por_subir.xlsx", 135 filas, 2026-09).

Construye objetos `Liquidacion` (mismo dataclass que devuelve
`procesar_empleado`) para poder reutilizar sin cambios
`_mapear_registro`/`_construir_conceptos`/`guardar_liquidacion` de
`core.repos.liquidaciones` -- ningún mapeo a columnas de Supabase se
duplica acá.

`parse_excel_liquidaciones` nunca lanza por datos de fila individual
malformados -- salta esa fila y la reporta en `errores`. No hace ninguna
conexión a SQL Server/Supabase (importable con pytest sin red).
"""

from __future__ import annotations

import datetime as dt
import io

import openpyxl

from core.repos.liquidaciones import (
    MESES_NOMBRE,
    DetalleMesDecimo,
    Liquidacion,
    periodos_decima_tercera,
)
from core.utils import normalizar_cedula

# Índices de columna (fila de encabezado = fila 2 del Excel; fila 1 viene
# vacía en el formato de origen). Verificado 2026-09 contra
# "Liquidaciones_por_subir.xlsx" (135 filas, 83 columnas).
COL = {
    "codigo": 0, "nombre_completo": 1, "cargo": 2, "sueldo_basico": 3,
    "puesto_servicio": 4, "cedula": 5, "motivo": 6, "seccion": 7,
    "fecha_ingreso": 8, "fecha_salida": 9, "dias": 10, "sueldo": 11,
    "reembolsos": 12, "fondo_reserva": 15, "bonificacion": 16,
    "maniobras": 17, "movilizacion": 18,
    "horas_25": 19, "horas_50": 20, "horas_100": 21,
    "sobt_25": 22, "sobt_50": 23, "sobt_100": 24,
    "dec13_ant": 26, "dec13_act": 27, "dec14_ant": 28, "dec14_act": 29,
    "vac_anteriores": 30, "vac_actuales": 31, "vac_pendientes": 32,
    "desahucio": 33, "indem_despido": 34,
    "otras_indem": 35, "valor_no_considerado": 36,
    "iess": 38, "prest_quirografario": 39, "prest_compania": 40,
    "anticipo_sueldo": 41, "anticipos_otros": 42, "anticipos_surtidos": 43,
    "iess_conyuge": 44, "impuesto_renta": 45, "multas": 46,
    "pension_alimenticia": 47, "prest_hipotecario": 48,
    "anticipos_otros_l": 49, "anticipo_l_desahucio": 50,
    "total_descuentos": 51, "total_a_recibir": 52,
}
PRIMERA_COL_MENSUAL = 63  # a partir de acá: fechas (día 1 de cada mes) en el encabezado


def _num(v) -> float:
    """'1.097,73' / ',00' / 132.30 / None -> float. Formato europeo (punto
    de miles, coma decimal) -- así vienen las celdas numéricas de este
    Excel (se guardaron como texto, no como número)."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return 0.0
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _cedula(v) -> str:
    if isinstance(v, (int, float)):
        return str(int(v)).zfill(10)
    return normalizar_cedula(str(v or "").strip())


def _fecha_iso(v) -> str:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    return ""


def _partir_nombre(nombre_completo: str) -> tuple[str, str]:
    """Heurística: convención ecuatoriana 'apellido apellido nombre nombre'
    -- se parte la lista de palabras a la mitad (primera mitad apellidos,
    segunda mitad nombres; con cantidad impar, la palabra extra queda del
    lado de los apellidos). NO es exacto para nombres compuestos
    irregulares -- si se necesita precisión, cruzar por cédula contra
    `rpemplea` del lado de quien llama (este parser no hace I/O a
    propósito, tiene que poder importarse con pytest sin red)."""
    palabras = (nombre_completo or "").split()
    if not palabras:
        return "", ""
    mitad = (len(palabras) + 1) // 2
    return " ".join(palabras[:mitad]), " ".join(palabras[mitad:])


def parse_excel_liquidaciones(datos: bytes) -> tuple[list[Liquidacion], list[str]]:
    """Parsea un Excel del formato LIQUIDACIONES_REG a una lista de
    `Liquidacion` con los montos YA CALCULADOS tal cual vienen en el
    archivo (no vuelve a correr `procesar_empleado`). El resultado es
    directamente compatible con `_mapear_registro`/`_construir_conceptos`/
    `guardar_liquidacion` de `core.repos.liquidaciones`.

    Limitaciones conocidas (documentadas, no resueltas en silencio):
    - `apellidos`/`nombres` se parten de la columna combinada "APELLIDOS Y
      NOMBRES" con una heurística mitad/mitad -- puede fallar en nombres
      compuestos irregulares. Sin I/O acá no se puede cruzar por cédula
      contra `rpemplea` para corregirlo; hacerlo es responsabilidad de
      quien llama, si lo necesita.
    - Las columnas "OTRAS INDEMNIZACIONES" y "POR CUALQUIER VALOR NO
      CONSIDERADO" se devuelven en `Liquidacion.campos["OTRAS_INDEM"]`/
      `["VALOR_NO_CONSIDERADO"]` -- mismas claves que usa
      `_CONCEPTOS_DETALLE` de `core.repos.liquidaciones` (agregado
      2026-09 para Editar en cuadrícula/Cuadre masivo), así que
      `_construir_conceptos` SÍ las va a incluir si vienen distintas de 0.
    - "DECIMA CUARTA REMUNERACION_ANTERIOR" viene con signo negativo en
      algunas filas del Excel de origen -- se pasa tal cual, sin
      "corregir" el signo (observado, no explicado).
    - El desglose mensual de Décima Tercera ACTUAL (para
      `liquidaciones_periodos_calculo`, que lee el bot MRL) se arma
      recortando las columnas mensuales dinámicas del Excel al periodo
      que calcula `periodos_decima_tercera(fecha_ingreso, fecha_salida)`
      -- si esas columnas no cubren el rango completo, el desglose queda
      incompleto para esos meses (no se inventa un valor). La ANTERIOR
      nunca se detalla, igual que en `procesar_empleado`.
    """
    liquidaciones: list[Liquidacion] = []
    errores: list[str] = []

    wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    if len(filas) < 3:
        return [], ["El archivo no tiene filas de datos (se esperaba encabezado en la fila 2)."]

    encabezado = filas[1]
    meses_col: dict[tuple[int, int], int] = {}
    for i in range(PRIMERA_COL_MENSUAL, len(encabezado)):
        v = encabezado[i]
        if isinstance(v, (dt.datetime, dt.date)):
            fecha = v.date() if isinstance(v, dt.datetime) else v
            meses_col[(fecha.year, fecha.month)] = i

    for n, fila in enumerate(filas[2:], start=3):
        try:
            cedula_raw = fila[COL["cedula"]]
            if cedula_raw in (None, ""):
                continue  # fila vacía / de relleno del template
            ced = _cedula(cedula_raw)
            if len(ced) < 6:
                errores.append(f"Fila {n}: cédula inválida ({cedula_raw!r}), se salta.")
                continue

            fecha_ing = _fecha_iso(fila[COL["fecha_ingreso"]])
            fecha_sal = _fecha_iso(fila[COL["fecha_salida"]])
            if not fecha_sal:
                errores.append(f"Fila {n} (cédula {ced}): sin fecha de salida, se salta.")
                continue

            nombre_completo = str(fila[COL["nombre_completo"]] or "").strip()
            apellidos, nombres = _partir_nombre(nombre_completo)

            motivo_raw = str(fila[COL["motivo"]] or "").strip().upper()
            motivo = "RENUNCIA VOLUNTARIA" if motivo_raw in ("", "FALTA") else motivo_raw

            dias_raw = fila[COL["dias"]]
            try:
                dias_trabajados = int(round(float(dias_raw))) if dias_raw is not None else 0
            except (TypeError, ValueError):
                dias_trabajados = 0

            campos = {
                "SUELDO": _num(fila[COL["sueldo"]]),
                "REEMBOLSOS": _num(fila[COL["reembolsos"]]),
                "FONDO_RESERVA": _num(fila[COL["fondo_reserva"]]),
                "BONIFICACION": _num(fila[COL["bonificacion"]]),
                "MANIOBRAS": _num(fila[COL["maniobras"]]),
                "MOVILIZACION": _num(fila[COL["movilizacion"]]),
                "HORAS_25": _num(fila[COL["horas_25"]]),
                "HORAS_50": _num(fila[COL["horas_50"]]),
                "HORAS_100": _num(fila[COL["horas_100"]]),
                "VAL_SOBT_25": _num(fila[COL["sobt_25"]]),
                "VAL_SOBT_50": _num(fila[COL["sobt_50"]]),
                "VAL_SOBT_100": _num(fila[COL["sobt_100"]]),
                "DECIMA_TERCERA_ANTERIOR": _num(fila[COL["dec13_ant"]]),
                "DECIMA_TERCERA_ACTUAL": _num(fila[COL["dec13_act"]]),
                "DECIMA_CUARTA_ANTERIOR": _num(fila[COL["dec14_ant"]]),
                "DECIMA_CUARTA_ACTUAL": _num(fila[COL["dec14_act"]]),
                "VACACIONES_ANTERIOR": _num(fila[COL["vac_anteriores"]]),
                "VACACIONES_ULTIMO": _num(fila[COL["vac_actuales"]]),
                "VACACIONES_CALCULADAS": _num(fila[COL["vac_pendientes"]]),
                "DESAHUCIO": _num(fila[COL["desahucio"]]),
                "INDEM_DESPIDO": _num(fila[COL["indem_despido"]]),
                "OTRAS_INDEM": _num(fila[COL["otras_indem"]]),
                "VALOR_NO_CONSIDERADO": _num(fila[COL["valor_no_considerado"]]),
                "APORT_IESS": _num(fila[COL["iess"]]),
                "PRESTAMOS_QUIROGRAFARIOS": _num(fila[COL["prest_quirografario"]]),
                "PRESTAMOS_COMPANIA": _num(fila[COL["prest_compania"]]),
                "ANTICIPO_SUELDO": _num(fila[COL["anticipo_sueldo"]]),
                "ANTICIPOS_OTROS": _num(fila[COL["anticipos_otros"]]),
                "ANTICIPOS_SURTIDOS": _num(fila[COL["anticipos_surtidos"]]),
                "APORT_IESS_CONYUGE": _num(fila[COL["iess_conyuge"]]),
                "IMPUESTO_RENTA": _num(fila[COL["impuesto_renta"]]),
                "MULTAS": _num(fila[COL["multas"]]),
                "PENSION_ALIMENTICIA": _num(fila[COL["pension_alimenticia"]]),
                "PRESTAMO_HIPOTECARIO": _num(fila[COL["prest_hipotecario"]]),
                "ANTICIPOS_OTROS_L": _num(fila[COL["anticipos_otros_l"]]),
                "ANTICIPO_L_DESAHUCIO": _num(fila[COL["anticipo_l_desahucio"]]),
            }
            total_descuentos = _num(fila[COL["total_descuentos"]])
            total_a_recibir = _num(fila[COL["total_a_recibir"]])
            campos["TOTAL_DESCUENTOS"] = total_descuentos
            campos["TOTAL_A_RECIBIR"] = total_a_recibir
            # total_ingresos = total_a_recibir + total_descuentos (despejado
            # de total_a_recibir = ingresos - descuentos) -- más robusto
            # que sumar TOTAL_NOMINA_PENDIENTE + TOTAL_BENEFICIOS_SOCIALES
            # del Excel por separado, que podría acumular redondeos propios.
            campos["TOTAL_INGRESOS"] = round(total_a_recibir + total_descuentos, 2)

            # Desglose mensual de Décima Tercera ACTUAL (liquidaciones_
            # periodos_calculo, tipo='DEC_TERCERA') -- mismo criterio que
            # procesar_empleado: solo el periodo actual, nunca el anterior.
            detalle_dec13: list[DetalleMesDecimo] = []
            if fecha_ing:
                try:
                    fing_d = dt.date.fromisoformat(fecha_ing)
                    fsal_d = dt.date.fromisoformat(fecha_sal)
                    p13 = periodos_decima_tercera(fing_d, fsal_d)
                    if p13:
                        i_act, f_act, _pagado = p13[-1]
                        y, m = i_act.year, i_act.month
                        while dt.date(y, m, 1) <= f_act:
                            idx = meses_col.get((y, m))
                            if idx is not None and idx < len(fila):
                                detalle_dec13.append(DetalleMesDecimo(
                                    label=f"{MESES_NOMBRE[m]} -{y}", valor=_num(fila[idx])))
                            m, y = (1, y + 1) if m == 12 else (m + 1, y)
                except ValueError:
                    pass  # fechas raras -- se deja sin desglose mensual, no bloquea la fila

            liquidaciones.append(Liquidacion(
                empleado=str(fila[COL["codigo"]] or "").strip(),
                nombre=nombre_completo,
                cedula=ced,
                cargo=str(fila[COL["cargo"]] or "").strip(),
                depto=str(fila[COL["puesto_servicio"]] or "").strip(),
                seccion=str(fila[COL["seccion"]] or "").strip(),
                sueldo_base=_num(fila[COL["sueldo_basico"]]),
                fecha_ingreso=fecha_ing,
                fecha_salida=fecha_sal,
                motivo_salida=motivo,
                dias_trabajados=dias_trabajados,
                campos=campos,
                apellidos=apellidos,
                nombres=nombres,
                detalle_decimo_tercera=detalle_dec13,
            ))
        except Exception as e:  # noqa: BLE001 - una fila mala no debe tumbar el lote entero
            errores.append(f"Fila {n}: error inesperado parseando ({e}), se salta.")

    return liquidaciones, errores
