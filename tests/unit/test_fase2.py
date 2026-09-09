"""Fase 2: helpers de préstamos/observaciones que no tocan BD + migración SQLite."""

import sqlite3

from core.migrations_legacy.sqlite_to_appdb import migrar
from core.repos import prestamos
from core.repos.observaciones import _falta, _fila_obs


def test_fila_obs_filtra_refers_vacios():
    r = {"empleado": "1012", "fecha_ven": "2026-06-01", "APELLIDOS": "PEREIRA", "NOMBRES": "JUAN"}
    fila = _fila_obs(r, ["texto A", "", None, "texto B", "  "])
    assert fila.empleado == "1012"
    assert fila.apellidos_nombres == "PEREIRA JUAN"
    assert fila.textos == ["texto A", "texto B"]


def test_fila_obs_slots7_conserva_los_siete_incluidos_vacios():
    fila = _fila_obs({"empleado": "1"}, ["a", "", None, "b", "", "", ""])
    assert fila.slots7 == ["a", "", "", "b", "", "", ""]
    assert fila.textos == ["a", "b"]


def test_reporte_html_varios_concatena_por_empleado(monkeypatch):
    from core.repos import observaciones as obs

    monkeypatch.setattr(obs, "datos_basicos_empleado",
                        lambda c, f: {"nombre": f"EMP {c}"})
    monkeypatch.setattr(obs, "observaciones", lambda c, f: [])
    monkeypatch.setattr(obs, "multas", lambda c, f: [])
    monkeypatch.setattr(obs, "faltas", lambda c, f, **k: [])
    html = obs.reporte_html_varios("supabase", ["1012", "1013"])
    assert html.count("page-break-after") == 2
    assert "EMP 1012" in html and "EMP 1013" in html


def test_reporte_html_observaciones():
    from core.repos import observaciones as obs

    html = obs.reporte_html(
        "1012", "PEREIRA JUAN",
        [{"fecha_ven": "2026-01-01", "texto": "Atrasos <x>"}],
        [{"fecha": "2026-02-01", "valor": 10.0, "concepto": "MULTA", "observ": "n"}],
        [{"periodo": "2026-01", "ausencias": 1, "faltas_justificadas": 0, "faltas_injustificadas": 1, "total": 2}],
    )
    assert "<title>Observaciones" in html and "PEREIRA JUAN" in html
    assert "&lt;x&gt;" in html  # escapado
    assert "MULTA" in html


def test_historial_observaciones_html():
    from core.repos import observaciones as obs

    filas = [
        {"fecha_ven": "2026-06-15", "textos": ["Llegó tarde", "Sin uniforme <x>"]},
        {"fecha_ven": "2026-05-02", "textos": ["Felicitación del cliente"]},
    ]
    html = obs.historial_observaciones_html("1012", "PEREIRA JUAN", filas)
    assert "<!DOCTYPE html>" in html
    assert "PEREIRA JUAN" in html and "1012" in html
    assert "2 registro(s)" in html
    assert "<li>Llegó tarde</li>" in html
    assert "&lt;x&gt;" in html  # escapado


def test_historial_observaciones_html_vacio_no_revienta():
    from core.repos import observaciones as obs

    html = obs.historial_observaciones_html("1012", "X", [])
    assert "Sin observaciones" in html and "0 registro(s)" in html


def _xlsx_obs(filas):
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_carga_masiva_observaciones():
    import datetime as dt

    from core.excel.parsers import parse_carga_masiva_observaciones

    data = _xlsx_obs(
        [
            ["EMPLEADO", "PERIODO", "TEXTO"],
            ["1012", "2026-06", "Llegó tarde"],
            [1013, dt.datetime(2026, 5, 1), "  Sin uniforme  "],
            ["", "", ""],                       # fila vacía -> ignorada
            ["1014", "2026-13", "mes malo"],     # periodo inválido -> error
            ["1015", "2026-07", ""],             # sin texto -> error
        ]
    )
    filas, errores = parse_carga_masiva_observaciones(data)
    assert filas == [
        {"empleado": "1012", "periodo": "2026-06", "texto": "Llegó tarde"},
        {"empleado": "1013", "periodo": "2026-05", "texto": "Sin uniforme"},
    ]
    assert len(errores) == 2


def test_parse_carga_masiva_observaciones_falta_columna():
    from core.excel.parsers import parse_carga_masiva_observaciones

    filas, errores = parse_carga_masiva_observaciones(_xlsx_obs([["EMPLEADO", "TEXTO"], ["1012", "x"]]))
    assert filas == [] and "PERIODO" in errores[0]


def test_job_carga_masiva_observaciones(monkeypatch, tmp_path):
    from core.repos import observaciones as obs

    llamadas = []

    def _fake_guardar(empleado, periodo, texto, *, usuario, roles):
        llamadas.append((empleado, periodo, texto))
        return "duplicado" if texto == "repe" else "refer1"

    monkeypatch.setattr(obs, "guardar_observacion", _fake_guardar)
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path))

    class _Ctx:
        job_id = 1
        cancelado = False

        def progreso(self, *a):
            pass

        def set_resultado(self, r):
            self.resultado = r

    ctx = _Ctx()
    obs.job_carga_masiva_observaciones(
        ctx,
        [
            {"empleado": "1012", "periodo": "2026-06", "texto": "ok"},
            {"empleado": "1013", "periodo": "2026-06", "texto": "repe"},
        ],
        usuario="admin",
        roles={"admin"},
    )
    assert len(llamadas) == 2
    assert ctx.resultado.endswith(".xlsx")


def test_falta_suma_total():
    f = _falta({"FECHA_VEN": "2026-06-15", "TOTAUS": 8, "TOTFJ": 0, "TOTFI": 4})
    assert f.periodo == "2026-06"
    assert f.total == 12.0


def test_numeros_migrados_se_excluyen_de_rphistor():
    assert "27958" in prestamos._NUMEROS_MIGRADOS
    assert prestamos.CLASE_PRESTAMO == 205


def test_agrupar_por_numero():
    M = prestamos.MovimientoPrestamo
    movs = [
        M("2025-02-05", 100.0, "CUOTA", "500", "RPHISTOR", tipo="pago"),
        M("2025-03-05", 100.0, "CUOTA", "500", "RPHISTOR", tipo="pago"),
        M("2025-06-30", 800.0, "SALDO", "500", "RPINGDES", tipo="pendiente"),
        M("2025-04-05", 300.0, "SALDO", "700", "RPINGDES", tipo="pendiente"),
    ]
    g = {r.numero: r for r in prestamos.agrupar_por_numero(movs)}
    assert g["500"].abonado == 200.0          # pagos de RPHISTOR
    assert g["500"].saldo == 800.0            # pendiente de RPINGDES
    assert g["500"].prestado == 1000.0        # sintético: 200 + 800
    assert g["500"].cuotas == 2
    assert g["500"].cuota_promedio == 100.0
    assert g["500"].cancelado is False
    assert g["500"].meses_brecha == 0  # feb -> mar consecutivos
    assert "para cancelar" in g["500"].estado
    assert g["700"].saldo == 300.0 and g["700"].abonado == 0.0


def test_agrupar_prestado_usa_desembolso_real_si_existe():
    M = prestamos.MovimientoPrestamo
    movs = [
        M("2020-01-01", 1500.0, "PRESTAMO", "MIG_1", "MIGRADO", tipo="desembolso"),
        M("2020-02-01", 300.0, "CUOTA", "MIG_1", "MIGRADO", tipo="pago"),
        M("2020-03-01", 300.0, "CUOTA", "MIG_1", "MIGRADO", tipo="pago"),
    ]
    r = prestamos.agrupar_por_numero(movs)[0]
    assert r.prestado == 1500.0   # el desembolso real, no 300+300
    assert r.abonado == 600.0
    assert r.saldo == 0.0         # no hay RPINGDES pendiente
    assert r.cancelado is True


def test_saldo_total_solo_suma_pendiente(monkeypatch):
    M = prestamos.MovimientoPrestamo
    fake = [
        M("2025-02-05", 100.0, "CUOTA", "9", "RPHISTOR", tipo="pago"),
        M("2025-03-05", 100.0, "CUOTA", "9", "RPHISTOR", tipo="pago"),
        M("2025-06-30", 250.0, "SALDO", "9", "RPINGDES", tipo="pendiente"),
    ]
    monkeypatch.setattr(prestamos, "historial_empleado", lambda c, f: fake)
    assert prestamos.saldo_total("9", "sqlserver") == 250.0  # no 450


def test_num_norm_normaliza_float_y_string():
    f = prestamos._num_norm
    assert f("35923") == "35923"
    assert f("35923.0") == "35923"      # NUMERO es float en las tablas -> str da '35923.0'
    assert f(35923.0) == "35923"
    assert f(35923) == "35923"
    assert f("  35923 ") == "35923"
    assert f(None) == ""
    assert f("MIG_5") == "MIG_5"        # no numérico -> tal cual
    # todos los números migrados normalizan a algo que está en el frozenset
    assert all(f(n + ".0") in prestamos._NUMEROS_MIGRADOS for n in prestamos._NUMEROS_MIGRADOS)


def test_agrupar_detecta_brecha_y_cancelado():
    M = prestamos.MovimientoPrestamo
    movs = [
        M("2025-01-31", 200.0, "C", "9", "RPHISTOR", tipo="pago"),
        M("2025-04-30", 200.0, "C", "9", "RPHISTOR", tipo="pago"),  # brecha feb y mar
    ]
    r = prestamos.agrupar_por_numero(movs)[0]
    assert r.saldo == 0.0 and r.cancelado is True  # sin RPINGDES pendiente
    assert r.abonado == 400.0
    assert r.meses_brecha == 2
    assert "Cancelado" in r.estado
    det = prestamos.movimientos_de_numero(movs, "9")
    assert len(det) == 2 and det[0].fecha == "2025-01-31"


def test_filtrar_movimientos():
    from dataclasses import asdict

    M = prestamos.MovimientoPrestamo
    movs = [
        asdict(M("2025-02-28", 100.0, "CUOTA", "9", "RPHISTOR", tipo="pago")),
        asdict(M("2025-03-31", 250.0, "CUOTA EXTRA", "12", "RPHISTOR", tipo="pago")),
        asdict(M("2025-06-30", 500.0, "SALDO", "9", "RPINGDES", tipo="pendiente")),
        asdict(M("2025-04-30", 50.0, "CUADRE", "9", "MIGRADO", tipo="pago", es_cuadre=True)),
    ]
    f = prestamos.filtrar_movimientos
    assert len(f(movs, tipo="pago")) == 3
    assert len(f(movs, tipo="pendiente")) == 1
    assert len(f(movs, tipo="egreso")) == 3   # alias viejo
    assert len(f(movs, tipo="ingreso")) == 1  # alias viejo
    assert [m["numero"] for m in f(movs, numero="12")] == ["12"]
    assert len(f(movs, origen="RPHISTOR")) == 2
    assert len(f(movs, texto="extra")) == 1
    assert len(f(movs, desde="2025-04-01")) == 2
    assert len(f(movs, monto_min=200)) == 2  # 250 y 500
    assert len(f(movs, monto_max=100)) == 2  # 100 y 50
    assert f(movs) == movs  # sin filtros, todo


def test_historial_xlsx_tiene_hoja_resumen():
    import io

    import openpyxl

    from core.excel.prestamos_builders import historial_xlsx

    M = prestamos.MovimientoPrestamo
    data = historial_xlsx("1012", "PEREIRA", [
        M("2025-01-05", 100.0, "CUOTA", "9", "RPHISTOR", tipo="pago"),
        M("2025-06-30", 400.0, "SALDO", "9", "RPINGDES", tipo="pendiente"),
    ])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert {"Historial", "Resumen por préstamo"} <= set(wb.sheetnames)
    ws = wb["Historial"]
    assert "HISTORIAL DE PRÉSTAMOS" in str(ws["A1"].value)
    valores = [c.value for row in ws.iter_rows() for c in row]
    assert "SALDO PENDIENTE" in valores
    assert 400.0 in valores  # el saldo pendiente, no la suma de todo


def test_migracion_sqlite_a_appdb(app_db, tmp_path):
    ruta = tmp_path / "Saldo_prestamos_driver.db"
    con = sqlite3.connect(ruta)
    con.execute(
        "CREATE TABLE historial_prestamos (codigo_empleado TEXT, fecha TEXT, "
        "ingreso REAL, egreso REAL, concepto TEXT, tipo TEXT, numero_fila INTEGER)"
    )
    con.executemany(
        "INSERT INTO historial_prestamos VALUES (?,?,?,?,?,?,?)",
        [
            ("1012", "2020-01-15", 0, 50.0, "PRESTAMO COMPANIA", "NORMAL", 1),
            ("1012", "2020-06-01", 100.0, 0, "CUADRE", "CUADRE", 2),
        ],
    )
    con.commit()
    con.close()

    n = migrar(str(ruta))
    assert n == 2

    movs = prestamos._historial_migrado("1012")
    assert len(movs) == 2
    assert any(m.es_cuadre for m in movs)
    assert any(m.valor == 50.0 and m.tipo == "pago" for m in movs)     # egreso migrado
    assert any(m.valor == 100.0 and m.tipo == "desembolso" for m in movs)  # ingreso migrado

    # idempotente: correr de nuevo no duplica
    assert migrar(str(ruta)) == 0
    assert len(prestamos._historial_migrado("1012")) == 2

    # una fila nueva en el SQLite sí entra
    con = sqlite3.connect(ruta)
    con.execute(
        "INSERT INTO historial_prestamos VALUES (?,?,?,?,?,?,?)",
        ("1012", "2021-03-10", 0, 25.0, "CUOTA", "NORMAL", 3),
    )
    con.commit()
    con.close()
    assert migrar(str(ruta)) == 1
    assert len(prestamos._historial_migrado("1012")) == 3

    # --reemplazar deja solo lo del SQLite
    assert migrar(str(ruta), reemplazar=True) == 3


def test_historial_display_reconstruye_filas_del_arbol(monkeypatch):
    """Una fila INGRESO sintética por préstamo + una EGRESO por pago, ordenadas
    por fecha, con saldo progresivo — como `buscar_prestamos` del .pyw."""
    M = prestamos.MovimientoPrestamo
    fake = [
        M("2025-01-02", 300.0, "SALDO", "9", "RPINGDES", tipo="pendiente"),  # fecha de registro
        M("2025-02-05", 100.0, "CUOTA FEB", "9", "RPHISTOR", tipo="pago"),
        M("2025-03-05", 100.0, "CUOTA MAR", "9", "RPHISTOR", tipo="pago"),
    ]
    monkeypatch.setattr(prestamos, "historial_empleado", lambda c, f: fake)
    monkeypatch.setattr(prestamos, "_datos_empleado_prestamos", lambda c, f: ("PEREZ JUAN", "0912345678"))

    crudas, info = prestamos.historial_display("9", "sqlserver")
    filas = prestamos.numerar_historial(crudas)

    # 1 ingreso (100+100+300 = 500) + 2 egresos
    assert [x.tipo for x in filas] == ["INGRESO", "EGRESO", "EGRESO"]
    ingreso = filas[0]
    assert ingreso.ingreso == 500.0 and ingreso.egreso == 0.0
    assert ingreso.saldo == 500.0
    assert "[H]" in ingreso.numero          # el grupo tiene movimientos históricos
    assert filas[1].egreso == 100.0 and filas[1].saldo == 400.0
    assert filas[2].saldo == 300.0          # queda el pendiente
    assert filas[0].fecha_fmt == "02/01/2025"
    assert info.nombre == "PEREZ JUAN"
    assert info.cedula == "0912345678"
    assert info.saldo_total == 300.0
    assert info.historicos == 2

    # numerar sobre un subconjunto recalcula # y saldo (igual que el .pyw)
    solo_egresos = prestamos.filtrar_historial(crudas, tipo="EGRESO")
    fe = prestamos.numerar_historial(solo_egresos)
    assert [x.posicion for x in fe] == [1, 2]
    assert fe[0].saldo == -100.0


def test_historial_display_usa_desembolso_real_migrado(monkeypatch):
    M = prestamos.MovimientoPrestamo
    fake = [
        M("2020-01-01", 1500.0, "PRESTAMO", "MIG_1", "MIGRADO", tipo="desembolso"),
        M("2020-02-01", 300.0, "CUOTA", "MIG_1", "MIGRADO", tipo="pago"),
    ]
    monkeypatch.setattr(prestamos, "historial_empleado", lambda c, f: fake)
    monkeypatch.setattr(prestamos, "_datos_empleado_prestamos", lambda c, f: ("", ""))
    crudas, _ = prestamos.historial_display("1", "sqlserver")
    filas = prestamos.numerar_historial(crudas)
    assert filas[0].tipo == "INGRESO" and filas[0].ingreso == 1500.0
    assert filas[0].historico is True
    assert filas[1].tipo == "EGRESO" and filas[1].egreso == 300.0
    assert filas[1].saldo == 1200.0


def test_filtrar_historial_por_origen_y_fecha():
    crudas = [
        {"fecha": "2025-01-10", "tipo": "INGRESO", "valor": 500.0, "numero": "9",
         "observacion": "PRESTAMO", "origen": "RPINGDES", "historico": False},
        {"fecha": "2025-02-10", "tipo": "EGRESO", "valor": 100.0, "numero": "9",
         "observacion": "CUOTA", "origen": "RPHISTOR", "historico": True},
    ]
    f = prestamos.filtrar_historial
    assert len(f(crudas, origen="SISTEMA")) == 1
    assert len(f(crudas, origen="HISTORICO")) == 1
    assert len(f(crudas, desde="2025-02-01")) == 1
    assert len(f(crudas, tipo="INGRESO")) == 1
    assert len(f(crudas, texto="cuota")) == 1
