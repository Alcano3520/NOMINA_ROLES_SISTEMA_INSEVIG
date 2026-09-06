"""Módulo admin: búsqueda de auditoría con filtros (BD de la app)."""

from __future__ import annotations

import datetime as dt

from core.db import appdb
from core.db.models import AuditLog
from core.repos.admin import buscar_auditoria


def _sembrar(filas: list[dict]) -> None:
    with appdb.session() as s:
        for f in filas:
            s.add(AuditLog(**f))
        s.commit()


def test_buscar_auditoria_filtra_por_modulo_estado_y_fecha(app_db):
    _sembrar(
        [
            {"username": "ana", "module": "empleados", "action": "editar", "status": "ok",
             "ts": dt.datetime(2026, 1, 10, 9, 0)},
            {"username": "beto", "module": "roles", "action": "generar", "status": "error",
             "ts": dt.datetime(2026, 3, 5, 12, 0)},
            {"username": "ana", "module": "empleados", "action": "eliminar", "status": "ok",
             "ts": dt.datetime(2026, 6, 1, 8, 0)},
        ]
    )

    total, filas = buscar_auditoria(modulo="empleados")
    assert total == 2 and len(filas) == 2
    assert filas[0]["ts"].startswith("2026-06-01")  # más reciente primero

    total, filas = buscar_auditoria(estado="error")
    assert total == 1 and filas[0]["usuario"] == "beto"

    total, filas = buscar_auditoria(desde="2026-02-01", hasta="2026-05-31")
    assert total == 1 and filas[0]["modulo"] == "roles"

    total, _ = buscar_auditoria(usuario="AN", desde="2026-05-01")  # substring, case-insensitive
    assert total == 1


def test_buscar_auditoria_hasta_es_inclusivo(app_db):
    _sembrar(
        [
            {"username": "u", "module": "m", "action": "a", "status": "ok",
             "ts": dt.datetime(2026, 4, 30, 23, 30)},
        ]
    )
    total, _ = buscar_auditoria(hasta="2026-04-30")
    assert total == 1


def test_buscar_auditoria_limite_no_afecta_al_total(app_db):
    _sembrar(
        [
            {"username": "u", "module": "m", "action": "a", "status": "ok",
             "ts": dt.datetime(2026, 1, 1) + dt.timedelta(hours=i)}
            for i in range(30)
        ]
    )
    total, filas = buscar_auditoria(limite=10)
    assert total == 30
    assert len(filas) == 10


def test_buscar_auditoria_fecha_invalida_se_ignora(app_db):
    _sembrar(
        [{"username": "u", "module": "m", "action": "a", "status": "ok", "ts": dt.datetime(2026, 1, 1)}]
    )
    total, _ = buscar_auditoria(desde="no-es-fecha")
    assert total == 1


def test_auditoria_xlsx_es_un_xlsx_valido():
    import io

    import openpyxl

    from core.excel.admin_builders import auditoria_xlsx

    filas = [
        {"ts": "2026-06-01 08:00:00", "usuario": "ana", "modulo": "empleados",
         "accion": "editar", "objetivo": "RPEMPLEA 1012", "status": "ok"},
    ]
    wb = openpyxl.load_workbook(io.BytesIO(auditoria_xlsx(filas)))
    ws = wb.active
    assert ws["A1"].value == "FECHA/HORA"
    assert ws["B2"].value == "ana"
    assert ws["F2"].value == "ok"


def test_auditoria_xlsx_sin_filas_no_revienta():
    import io

    import openpyxl

    from core.excel.admin_builders import auditoria_xlsx

    wb = openpyxl.load_workbook(io.BytesIO(auditoria_xlsx([])))
    assert wb.active["A1"].value == "FECHA/HORA"
