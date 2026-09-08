"""Builders de Excel del módulo Vacaciones."""

from __future__ import annotations

import io

import openpyxl

from core.excel.vacaciones_builders import pendientes_global_xlsx, reporte_completo_xlsx

_FILAS_COMPLETO = [
    {
        "cedula": "0920116811", "apellidos": "PEREZ", "nombres": "JUAN", "cargo": "GUARDIA",
        "departamento": "OPERACIONES", "periodo": "2024-2025", "tipo": "pagada",
        "estado_doc": "completado", "desde": "2024-01-01", "hasta": "2024-12-31",
        "dias_tomados": 15, "dias_adicionales": 1, "total_pagar": 256.0, "anticipo": 0.0,
        "fecha_pago": "2025-02-10", "banco": "PICHINCHA", "no_cheque": "123", "observaciones": "",
    },
    {
        "cedula": "0911111111", "apellidos": "GOMEZ", "nombres": "ANA", "cargo": "SUPERVISOR",
        "departamento": "ADMIN", "periodo": "2024-2025", "tipo": "gozada",
        "estado_doc": "pendiente", "desde": "2024-06-01", "hasta": "2024-06-15",
        "dias_tomados": 15, "dias_adicionales": 0, "total_pagar": 0.0, "anticipo": 0.0,
        "fecha_pago": "", "banco": "", "no_cheque": "", "observaciones": "x",
    },
]

_FILAS_PENDIENTES = [
    {
        "cedula": "0920116811", "apellidos": "PEREZ", "nombres": "JUAN", "cargo": "GUARDIA",
        "departamento": "OPERACIONES", "fecha_ingreso": "2020-03-15", "periodo": "2023-2024",
        "dias_derecho": 16, "dias_adicionales": 1, "dias_gozados": 0, "dias_pendientes": 16,
    },
    {
        "cedula": "0911111111", "apellidos": "GOMEZ", "nombres": "ANA", "cargo": "SUPERVISOR",
        "departamento": "ADMIN", "fecha_ingreso": "2019-01-10", "periodo": "2023-2024",
        "dias_derecho": 17, "dias_adicionales": 2, "dias_gozados": 5, "dias_pendientes": 12,
    },
]


def test_reporte_completo_xlsx_es_valido():
    data = reporte_completo_xlsx(_FILAS_COMPLETO)
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    txt = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "REPORTE COMPLETO" in txt
    assert "0920116811" in txt and "PEREZ" in txt
    assert "GOCE" in txt and "PAGO" in txt  # tipo traducido


def test_pendientes_global_xlsx_suma_total_y_columnas():
    data = pendientes_global_xlsx(_FILAS_PENDIENTES)
    assert data[:2] == b"PK"
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    txt = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "VACACIONES PENDIENTES (GLOBAL)" in txt
    assert "28 días pendientes en total" in txt  # 16 + 12
    assert "Días Pendientes" in txt
    assert "0911111111" in txt


def test_pendientes_global_xlsx_vacio_no_rompe():
    data = pendientes_global_xlsx([])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["A4"].value == "Cédula"
